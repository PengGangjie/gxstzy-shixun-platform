# -*- coding: utf-8 -*-
"""实训科管理平台 · FastAPI（Logto + Turso + 静态站点 + 教务处权限后台）。"""
from __future__ import annotations

import logging
import os
import time
from collections import defaultdict, deque
from typing import Any, Union

from fastapi import FastAPI, File, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from logto import LogtoClient, LogtoConfig, Storage, UserInfoScope
from pydantic import BaseModel, Field
from starlette.middleware.sessions import SessionMiddleware

from .config import get_settings
from .db import get_user_by_sub, list_users, ping_db, set_user_role, upsert_user
from .roles import (
    COLLEGES,
    ROLE_LABELS,
    ROLES,
    can_assign_role,
    capabilities,
    normalize_role,
    public_user,
)
from . import cas_auth
from . import room_store

logger = logging.getLogger("shixun")

# 写操作限速：每用户 60 秒内最多 30 次（单实例内存窗口，防脚本刷写撑爆 Turso 配额）
WRITE_RATE_MAX = 30
WRITE_RATE_WINDOW = 60.0
_write_hits: dict[str, deque[float]] = defaultdict(deque)

# 台账导入上传体上限（xlsx 解压后按行数与字段截断，此处再限原始体积）
MAX_IMPORT_BYTES = 3 * 1024 * 1024


def _rate_limited(key: str) -> bool:
    now = time.monotonic()
    hits = _write_hits[key]
    while hits and now - hits[0] > WRITE_RATE_WINDOW:
        hits.popleft()
    if len(hits) >= WRITE_RATE_MAX:
        return True
    hits.append(now)
    if len(_write_hits) > 4096:
        for k in [k for k, v in _write_hits.items() if not v or now - v[-1] > WRITE_RATE_WINDOW]:
            _write_hits.pop(k, None)
    return False


def _write_limit_response(user: dict[str, Any] | None) -> JSONResponse | None:
    key = (user or {}).get("logto_sub") or "anon"
    if _rate_limited(f"w:{key}"):
        return JSONResponse({"detail": "操作过于频繁，请稍后再试"}, status_code=429)
    return None


def _safe_next(nxt: str | None) -> str | None:
    """登录后跳转目标仅允许站内路径；浏览器把 /\\ 开头视同 //，须一并拒绝。"""
    if not nxt or not nxt.startswith("/"):
        return None
    if nxt.startswith("//") or nxt.startswith("/\\"):
        return None
    return nxt


def _db_fail(prefix: str, exc: Exception) -> JSONResponse:
    raw = str(exc)
    if "quota" in raw.lower():
        return JSONResponse(
            {"detail": f"{prefix}云端数据库配额已满，请先删除部分教室照片后再试"},
            status_code=507,
        )
    # 原始异常可能含连接串/SQL 等内部信息，只进日志不回显
    logger.exception("%s（内部错误）", prefix)
    return JSONResponse({"detail": f"{prefix}服务内部错误，请稍后重试或联系管理员"}, status_code=500)

PUBLIC_PREFIXES = (
    "/health",
    "/sign-in",
    "/callback",
    "/sign-out",
    "/cas/login",
    "/cas/callback",
    "/assets/",
    "/brand/",
    "/icons/",
)

# 游客可浏览首页外壳；模块页 / 搜索索引 / 台账 JSON 仍须登录
GUEST_HOME_PATHS = frozenset(
    {
        "/",
        "/index.html",
        "/广西生态工程职业技术学院-教务处-实训科管理平台.html",
    }
)

# 静态资源扩展名：须已登录（/assets/ 已在 PUBLIC_PREFIXES）
STATIC_ASSET_SUFFIXES = (
    ".css",
    ".js",
    ".png",
    ".jpg",
    ".jpeg",
    ".gif",
    ".webp",
    ".svg",
    ".ico",
    ".woff",
    ".woff2",
    ".map",
    ".pdf",
)

settings = get_settings()
app = FastAPI(title="广西生态工程职业技术学院 · 实训科管理平台")


class SessionStorage(Storage):
    def __init__(self, session: dict) -> None:
        self._session = session

    def get(self, key: str) -> Union[str, None]:
        val = self._session.get(key)
        return None if val is None else str(val)

    def set(self, key: str, value: Union[str, None]) -> None:
        if value is None:
            self._session.pop(key, None)
        else:
            self._session[key] = value

    def delete(self, key: str) -> None:
        self._session.pop(key, None)


def logto_client(request: Request) -> LogtoClient:
    return LogtoClient(
        LogtoConfig(
            endpoint=settings.logto_endpoint,
            appId=settings.logto_app_id,
            appSecret=settings.logto_app_secret,
            scopes=[UserInfoScope.email, UserInfoScope.phone],
        ),
        storage=SessionStorage(request.session),
    )


def auth_configured() -> bool:
    return bool(settings.logto_endpoint and settings.logto_app_id and settings.logto_app_secret)


def _email_is_bootstrap_admin(email: str | None) -> bool:
    if not email or not settings.admin_emails:
        return False
    return email.strip().lower() in settings.admin_emails


def current_identity(request: Request) -> dict[str, Any] | None:
    """当前会话身份：CAS（本地会话）或 Logto（OIDC id_token）任一渠道已认证即有效。

    返回 {sub, username, email, name, source}；sub 为数据库 users.logto_sub 主键
    （CAS 用户为 cas:{工号}）。
    """
    cas = request.session.get("cas_user")
    if isinstance(cas, dict) and cas.get("sub"):
        return cas
    if auth_configured():
        client = logto_client(request)
        if client.isAuthenticated():
            claims = client.getIdTokenClaims()
            if claims and claims.sub:
                return {
                    "sub": claims.sub,
                    "username": None,
                    "email": getattr(claims, "email", None),
                    "name": getattr(claims, "name", None),
                    "source": "logto",
                }
    return None


def load_db_user(request: Request) -> dict[str, Any] | None:
    identity = current_identity(request)
    if not identity:
        return None
    return get_user_by_sub(identity["sub"])


def require_capability(user: dict[str, Any] | None, cap: str) -> JSONResponse | None:
    if not user:
        return JSONResponse({"detail": "未登录"}, status_code=401)
    if cap not in capabilities(user):
        return JSONResponse({"detail": "无权限"}, status_code=403)
    return None


class SetRoleBody(BaseModel):
    role: str
    college: str | None = None
    lab_rooms: list[str] = Field(default_factory=list)


def _unauth_response(request: Request, path: str):
    """未登录响应：API/静态资源用 401，页面跳转登录。"""
    if path.startswith("/api/") or path.endswith(STATIC_ASSET_SUFFIXES) or path.endswith(".json"):
        return JSONResponse({"detail": "未登录", "sign_in": "/sign-in"}, status_code=401)
    next_q = path if path.startswith("/") else f"/{path}"
    return RedirectResponse(f"/sign-in?next={next_q}")


@app.middleware("http")
async def security_headers(request: Request, call_next):
    resp = await call_next(request)
    resp.headers.setdefault("X-Content-Type-Options", "nosniff")
    resp.headers.setdefault("X-Frame-Options", "SAMEORIGIN")
    resp.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    resp.headers.setdefault("Permissions-Policy", "camera=(self), microphone=(), geolocation=(), payment=()")
    # Koyeb 反代后 scheme 常为 http，须看 x-forwarded-proto
    if request.url.scheme == "https" or request.headers.get("x-forwarded-proto") == "https":
        resp.headers.setdefault("Strict-Transport-Security", "max-age=15552000")
    return resp


@app.middleware("http")
async def require_auth(request: Request, call_next):
    if not settings.auth_required or not auth_configured():
        return await call_next(request)
    path = request.url.path
    if path.startswith(PUBLIC_PREFIXES):
        return await call_next(request)

    authenticated = current_identity(request) is not None

    # 游客：首页外壳 + /api/me（返回未登录状态）
    if not authenticated and (path in GUEST_HOME_PATHS or path == "/api/me"):
        return await call_next(request)

    if path.endswith(STATIC_ASSET_SUFFIXES) or path.endswith(".json"):
        if authenticated:
            return await call_next(request)
        return _unauth_response(request, path)

    if not authenticated:
        return _unauth_response(request, path)

    # 教务处后台页：仅 jw_admin；管理 API：jw_admin 或学院管理员（本院）
    if path == "/admin" or path.startswith("/admin/"):
        user = load_db_user(request)
        if not user or "admin.panel" not in capabilities(user):
            return JSONResponse(
                {
                    "detail": "需要教务处管理员权限",
                    "hint": "请联系已授权管理员，或将 ADMIN_EMAILS 设为您的邮箱后重新登录。",
                    "home": "/",
                },
                status_code=403,
            )
    elif path == "/api/admin" or path.startswith("/api/admin/"):
        user = load_db_user(request)
        caps = capabilities(user) if user else set()
        if "admin.panel" not in caps and "users.manage_college" not in caps:
            return JSONResponse({"detail": "需要管理员权限"}, status_code=403)

    return await call_next(request)


# SessionMiddleware 须后注册（更靠外层），保证 require_auth 内可读写 request.session
app.add_middleware(
    SessionMiddleware,
    secret_key=settings.session_secret,
    session_cookie="shixun_session",
    https_only=settings.logto_redirect_uri.startswith("https://"),
    same_site="lax",
    max_age=14 * 24 * 3600,
)


@app.get("/health")
async def health():
    body: dict[str, Any] = {"status": "ok", "app": settings.app_name}
    try:
        body["db_schema_version"] = ping_db()
    except Exception:  # noqa: BLE001
        # /health 公开，数据库异常详情只进日志，避免泄漏连接信息
        logger.exception("healthcheck: 数据库不可达")
        body["status"] = "degraded"
        body["db"] = "unreachable"
    return body


@app.get("/sign-in")
async def sign_in(request: Request):
    if not auth_configured():
        return RedirectResponse("/")
    nxt = _safe_next(request.query_params.get("next"))
    if nxt and nxt not in {"/sign-in", "/callback"}:
        request.session["post_login_next"] = nxt
    client = logto_client(request)
    url = await client.signIn(redirectUri=settings.logto_redirect_uri)
    return RedirectResponse(url)


@app.get("/callback")
async def callback(request: Request):
    """OIDC 回调。无 code/state 或会话丢失时返回说明页，避免裸 500。"""
    if not request.query_params.get("code"):
        return RedirectResponse("/sign-in")
    client = logto_client(request)
    try:
        await client.handleSignInCallback(str(request.url))
    except Exception as exc:  # noqa: BLE001
        logger.exception("登录回调失败（state/code 校验）")
        return JSONResponse(
            {
                "detail": "登录回调失败",
                "reason": f"OIDC 回调校验未通过（{type(exc).__name__}，详情见服务端日志）",
                "hint": "请从平台首页重新登录。若出现 invalid_client，请在 Logto 控制台核对 App ID/Secret 后重新部署。",
                "sign_in": "/sign-in",
            },
            status_code=400,
        )
    claims = client.getIdTokenClaims()
    if claims and claims.sub:
        email = getattr(claims, "email", None)
        upsert_user(
            claims.sub,
            email,
            getattr(claims, "name", None),
            getattr(claims, "phone_number", None),
            default_role="student",
            promote_to_jw_admin=_email_is_bootstrap_admin(email),
        )
    nxt = request.session.pop("post_login_next", None)
    nxt = _safe_next(nxt) or "/"
    return RedirectResponse(nxt)

@app.get("/sign-out")
async def sign_out(request: Request):
    was_cas = isinstance(request.session.get("cas_user"), dict)
    if was_cas and cas_auth.cas_enabled():
        # CAS 用户：清本地会话后引导认证平台注销 SSO 会话（公共机房防残留）
        request.session.clear()
        from urllib.parse import quote

        base = settings.logto_redirect_uri.rsplit("/", 1)[0] + "/"
        return RedirectResponse(f"{cas_auth.cas_logout_url()}?service={quote(base, safe='')}")
    client = logto_client(request)
    url = await client.signOut(postLogoutRedirectUri=settings.logto_post_logout_uri)
    request.session.clear()
    return RedirectResponse(url)


# ---- 学校统一身份认证（CAS）单点登录 ----


@app.get("/cas/login")
async def cas_login(request: Request):
    if not cas_auth.cas_enabled():
        return JSONResponse({"detail": "学校统一身份认证未配置（待信息中心提供认证地址）"}, status_code=404)
    nxt = _safe_next(request.query_params.get("next"))
    if nxt:
        request.session["post_login_next"] = nxt
    from urllib.parse import quote

    return RedirectResponse(
        f"{cas_auth.cas_login_url()}?service={quote(cas_auth.service_url(), safe='')}"
    )


@app.get("/cas/callback")
async def cas_callback(request: Request):
    if not cas_auth.cas_enabled():
        return JSONResponse({"detail": "学校统一身份认证未配置"}, status_code=404)
    ip = request.client.host if request.client else "?"
    if _rate_limited(f"cas:{ip}"):
        return JSONResponse({"detail": "尝试过于频繁，请稍后再试"}, status_code=429)
    ticket = (request.query_params.get("ticket") or "").strip()
    if not ticket:
        return RedirectResponse("/cas/login")
    username = cas_auth.validate_ticket(ticket)
    if not username:
        return JSONResponse(
            {
                "detail": "统一身份认证校验未通过。常见原因：① 应用回调地址尚未在认证平台注册；② 该账号未获访问授权；③ 票据已过期，请从首页重新登录。",
                "home": "/",
            },
            status_code=401,
        )
    request.session["cas_user"] = cas_auth.build_session_user(username)
    upsert_user(
        f"cas:{username}",
        None,
        username,
        None,
        employee_no=username,
        default_role="student",
        promote_to_jw_admin=username.strip().lower() in settings.admin_cas_accounts,
    )
    nxt = _safe_next(request.session.pop("post_login_next", None)) or "/"
    return RedirectResponse(nxt)


@app.get("/api/me")
async def me(request: Request):
    identity = current_identity(request)
    if not identity:
        return {
            "authenticated": False,
            # 前端据此展示「统一身份认证登录」入口（未配置时不显示）
            "cas_login": "/cas/login" if cas_auth.cas_enabled() else None,
        }
    sub = identity["sub"]
    user = get_user_by_sub(sub)
    if not user:
        email = identity.get("email")
        user = upsert_user(
            sub,
            email,
            identity.get("name"),
            None,
            employee_no=identity.get("username") if identity.get("source") == "cas" else None,
            default_role="student",
            promote_to_jw_admin=_email_is_bootstrap_admin(email)
            or (
                identity.get("source") == "cas"
                and str(identity.get("username") or "").strip().lower()
                in settings.admin_cas_accounts
            ),
        )
    body: dict[str, Any] = {
        "authenticated": True,
        "sub": sub,
        "source": identity.get("source"),
        "username": identity.get("username"),
        "employee_no": (user or {}).get("employee_no"),
        "email": identity.get("email"),
        "name": identity.get("name"),
        "cas_login": "/cas/login" if cas_auth.cas_enabled() else None,
    }
    if user:
        body["phone"] = user.get("phone")
        body.update(public_user(user))
        body["admin_panel"] = "/admin/" if "admin.panel" in capabilities(user) else None
    return body


@app.get("/api/meta/roles")
async def meta_roles(request: Request):
    user = load_db_user(request)
    err = require_capability(user, "admin.panel")
    if err:
        # 学院管理员也需要角色列表时放开只读元数据：已登录即可
        if not user:
            return err
    return {
        "roles": [{"id": r, "label": ROLE_LABELS[r]} for r in ROLES],
        "colleges": list(COLLEGES),
    }


@app.get("/api/admin/users")
async def admin_list_users(request: Request):
    user = load_db_user(request)
    err = require_capability(user, "users.manage_all")
    if err:
        # 学院管理员：本院列表
        err2 = require_capability(user, "users.manage_college")
        if err2:
            return err2
        college = (user or {}).get("college")
        if not college:
            return JSONResponse({"detail": "请先完善所属学院"}, status_code=400)
        rows = list_users(college=college)
        return {"users": [public_user(u) for u in rows], "scope": "college"}
    rows = list_users()
    return {"users": [public_user(u) for u in rows], "scope": "all"}


@app.post("/api/admin/users/{sub}/role")
async def admin_set_role(sub: str, body: SetRoleBody, request: Request):
    actor = load_db_user(request)
    if not actor:
        return JSONResponse({"detail": "未登录"}, status_code=401)
    err = _write_limit_response(actor)
    if err:
        return err
    if "users.manage_all" not in capabilities(actor) and "users.manage_college" not in capabilities(
        actor
    ):
        return JSONResponse({"detail": "无权限"}, status_code=403)

    role = normalize_role(body.role)
    college = body.college
    ok, msg = can_assign_role(actor, role, college)
    if not ok:
        return JSONResponse({"detail": msg}, status_code=403)

    # 教务处管理员可不填学院；其他角色建议填写
    if role != "jw_admin" and not (college or "").strip():
        return JSONResponse({"detail": "请选择所属学院"}, status_code=400)
    if college and college not in COLLEGES and role != "jw_admin":
        return JSONResponse({"detail": f"学院须为：{', '.join(COLLEGES)}"}, status_code=400)

    try:
        updated = set_user_role(
            sub,
            role,
            None if role == "jw_admin" else college,
            body.lab_rooms if role == "lab_tech" else [],
            actor_sub=actor.get("logto_sub"),
        )
    except KeyError as exc:
        return JSONResponse({"detail": str(exc)}, status_code=404)
    except Exception as exc:  # noqa: BLE001
        return JSONResponse({"detail": f"更新失败：{exc}"}, status_code=500)

    return {"ok": True, "user": public_user(updated)}


class RoomOverridesBody(BaseModel):
    overrides: dict[str, Any] = Field(default_factory=dict)


class RoomPhotoBody(BaseModel):
    data_url: str
    caption: str | None = None


class RoomEquipImportBody(BaseModel):
    rows: list[dict[str, Any]] = Field(default_factory=list)


def _actor_label(user: dict[str, Any] | None) -> str | None:
    if not user:
        return None
    return user.get("email") or user.get("name") or user.get("logto_sub")


def _parse_equip_upload(filename: str, raw: bytes) -> list[dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = raw.decode("utf-8-sig", errors="replace")
        lines = [ln for ln in text.splitlines() if ln.strip()]
        if not lines:
            return []
        import csv
        from io import StringIO

        reader = csv.DictReader(StringIO("\n".join(lines)))
        out: list[dict[str, Any]] = []
        for row in reader:
            out.append(dict(row))
            if len(out) >= room_store.MAX_EQUIP_ROWS:
                break
        return out
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(rows_iter, [])]
        out: list[dict[str, Any]] = []
        for row in rows_iter:
            item = {
                header[i]: ("" if row[i] is None else str(row[i]).strip())
                for i in range(min(len(header), len(row)))
                if header[i]
            }
            if any(item.values()):
                out.append(item)
            # 流式截断，防止超大工作簿撑爆内存
            if len(out) >= room_store.MAX_EQUIP_ROWS:
                break
        return out
    raise ValueError("请上传 .xlsx 或 .csv 文件")


@app.get("/api/rooms/{room_id}/state")
async def room_state(room_id: str, request: Request):
    user = load_db_user(request)
    if not user:
        return JSONResponse({"detail": "未登录"}, status_code=401)
    try:
        return room_store.get_room_state(room_id)
    except Exception as exc:  # noqa: BLE001
        return _db_fail("读取失败：", exc)


@app.put("/api/rooms/{room_id}/overrides")
async def room_save_overrides(room_id: str, body: RoomOverridesBody, request: Request):
    user = load_db_user(request)
    err = require_capability(user, "rooms.write")
    if err:
        return err
    err = _write_limit_response(user)
    if err:
        return err
    try:
        return room_store.save_overrides(room_id, body.overrides, _actor_label(user))
    except ValueError as exc:
        return JSONResponse({"detail": str(exc)}, status_code=400)
    except Exception as exc:  # noqa: BLE001
        return _db_fail("保存失败：", exc)


@app.post("/api/rooms/{room_id}/photos")
async def room_add_photo(room_id: str, body: RoomPhotoBody, request: Request):
    user = load_db_user(request)
    err = require_capability(user, "rooms.write")
    if err:
        return err
    err = _write_limit_response(user)
    if err:
        return err
    try:
        photo = room_store.add_photo(
            room_id, body.data_url, body.caption, _actor_label(user)
        )
        return {"ok": True, "photo": photo}
    except ValueError as exc:
        return JSONResponse({"detail": str(exc)}, status_code=400)
    except Exception as exc:  # noqa: BLE001
        return _db_fail("上传失败：", exc)


@app.delete("/api/rooms/{room_id}/photos/{photo_id}")
async def room_del_photo(room_id: str, photo_id: int, request: Request):
    user = load_db_user(request)
    err = require_capability(user, "rooms.write")
    if err:
        return err
    err = _write_limit_response(user)
    if err:
        return err
    try:
        room_store.delete_photo(room_id, photo_id, _actor_label(user))
        return {"ok": True}
    except Exception as exc:  # noqa: BLE001
        return _db_fail("删除失败：", exc)


@app.post("/api/rooms/{room_id}/equipment")
async def room_set_equipment(room_id: str, body: RoomEquipImportBody, request: Request):
    user = load_db_user(request)
    err = require_capability(user, "rooms.write")
    if err:
        return err
    err = _write_limit_response(user)
    if err:
        return err
    try:
        equipment = room_store.replace_equipment(room_id, body.rows or [], _actor_label(user))
        return {"ok": True, "count": len(equipment), "equipment": equipment}
    except ValueError as exc:
        return JSONResponse({"detail": str(exc)}, status_code=400)
    except Exception as exc:  # noqa: BLE001
        return _db_fail("保存失败：", exc)


@app.post("/api/rooms/{room_id}/equipment/import")
async def room_import_equipment(room_id: str, request: Request, file: UploadFile = File(...)):
    user = load_db_user(request)
    err = require_capability(user, "rooms.write")
    if err:
        return err
    err = _write_limit_response(user)
    if err:
        return err
    try:
        raw = await file.read()
        if len(raw) > MAX_IMPORT_BYTES:
            return JSONResponse(
                {"detail": f"文件过大（上限 {MAX_IMPORT_BYTES // 1024 // 1024}MB）"},
                status_code=413,
            )
        rows = _parse_equip_upload(file.filename or "equip.csv", raw)
        equipment = room_store.replace_equipment(room_id, rows, _actor_label(user))
        return {"ok": True, "count": len(equipment), "equipment": equipment}
    except ValueError as exc:
        return JSONResponse({"detail": str(exc)}, status_code=400)
    except Exception as exc:  # noqa: BLE001
        return _db_fail("导入失败：", exc)


@app.get("/admin")
@app.get("/admin/")
async def admin_panel():
    """教务处权限后台（中间件已校验 jw_admin）。"""
    page = settings.static_dir / "admin" / "index.html"
    if not page.is_file():
        return JSONResponse({"detail": "后台页面缺失"}, status_code=503)
    return FileResponse(page)


static_dir = settings.static_dir
if static_dir.is_dir():
    app.mount("/", StaticFiles(directory=str(static_dir), html=True), name="static")
else:

    @app.get("/")
    async def missing_static():
        return JSONResponse(
            {"detail": f"static 目录不存在: {static_dir}，请先运行 sync_shixun_platform_static.py"},
            status_code=503,
        )


def main() -> None:
    import uvicorn

    port = int(os.getenv("PORT", settings.port))
    uvicorn.run("app.main:app", host="0.0.0.0", port=port, reload=False)


if __name__ == "__main__":
    main()
